import os
import pickle
import json
from flask import (Flask, request, render_template, session, redirect, url_for, flash, send_file, render_template_string, jsonify)
from werkzeug.utils import secure_filename
import pandas as pd
from io import BytesIO
import uuid
from functools import wraps
import secrets
import string
import shutil
import smtplib
from email.message import EmailMessage
from dotenv import load_dotenv
from itsdangerous import URLSafeTimedSerializer, SignatureExpired, BadSignature
from extensions import db
from models.user import User
from models.turma import Turma
from models.grupo import Grupo
from models.aluno import Aluno
from services.save_studant import save_studant

# Carregar variáveis de ambiente do arquivo .env
load_dotenv()

import re

from services.pdf_parser import processar_pdfs_para_dataframe
from services.group_generator import formar_grupos, horarios_conflitam
from services.group_manager import sugerir_trocas, executar_troca
from services.analytics_manager import calcular_ranking_disponibilidade

# Grade fixa da planilha: 07:00 → 18:00, de 30 em 30 min (igual ao frontend)
HORARIOS_GRADE = [
    "07:00", "07:30", "08:00", "08:30", "09:00", "09:30",
    "10:00", "10:30", "11:00", "11:30", "12:00", "12:30",
    "13:00", "13:30", "14:00", "14:30", "15:00", "15:30",
    "16:00", "16:30", "17:00", "17:30", "18:00",
]

def _hhmm_to_minutes(s: str) -> int:
    hh, mm = s.strip().split(":", 1)
    return int(hh) * 60 + int(mm)

def _minutes_to_hhmm(total: int) -> str:
    hh, mm = divmod(total, 60)
    return f"{hh:02d}:{mm:02d}"

def obter_intervalo_minutos(horario_val):
    if pd.isna(horario_val):
        return None
    try:
        val = str(horario_val).strip()
        if " - " in val:
            a, b = val.split(" - ", 1)
        elif "-" in val:
            a, b = val.split("-", 1)
        else:
            ini = _hhmm_to_minutes(val) if ":" in val else int(float(val) * 60)
            return (ini, ini + 30)
        return (_hhmm_to_minutes(a), _hhmm_to_minutes(b))
    except (ValueError, TypeError, IndexError):
        return None

def obter_hora_inicio(horario_val):
    intervalo = obter_intervalo_minutos(horario_val)
    if intervalo is None:
        return None
    return _minutes_to_hhmm(intervalo[0])

def format_horario_slot(slot: str) -> str:
    ini = _hhmm_to_minutes(slot)
    return f"{slot} - {_minutes_to_hhmm(ini + 30)}"

def slots_ocupados_no_intervalo(ini_min: int, fim_min: int):
    for slot in HORARIOS_GRADE:
        slot_min = _hhmm_to_minutes(slot)
        if ini_min < slot_min + 30 and fim_min > slot_min:
            yield slot


from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
from werkzeug.security import check_password_hash, generate_password_hash

template_dir = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__, 
            template_folder=os.path.join(template_dir, 'templates'),
            static_folder=os.path.join(template_dir, 'static'))

app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'ACADEMYSYNC_KEY')
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['SESSION_DATA_FOLDER'] = 'session_data'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['SESSION_DATA_FOLDER'], exist_ok=True)

# --- Configuração do Banco de Dados PostgreSQL/SQLite ---
database_url = os.environ.get('DATABASE_URL')
if database_url and database_url.startswith("postgres://"):
    # SQLAlchemy exige 'postgresql://' em vez de 'postgres://'
    database_url = database_url.replace("postgres://", "postgresql://", 1)

app.config['SQLALCHEMY_DATABASE_URI'] = database_url or 'sqlite:///usuarios.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

# Inicializa as tabelas do banco de dados se não existirem
with app.app_context():
    db.create_all()
    # Migrações simples de colunas para compatibilidade
    try:
        db.session.execute(db.text('ALTER TABLE users ADD COLUMN security_question VARCHAR(256)'))
        db.session.commit()
    except Exception:
        db.session.rollback()
        
    try:
        db.session.execute(db.text('ALTER TABLE users ADD COLUMN security_answer_hash VARCHAR(256)'))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text('ALTER TABLE turmas ADD COLUMN em_imersao BOOLEAN DEFAULT FALSE'))
        db.session.commit()
    except Exception:
        db.session.rollback()

    try:
        db.session.execute(db.text('ALTER TABLE alunos ADD COLUMN matricula VARCHAR(50)'))
        db.session.commit()
    except Exception:
        db.session.rollback()

    # Alunos podem existir antes de serem alocados a um grupo
    try:
        db.session.execute(db.text('ALTER TABLE alunos ALTER COLUMN grupo_id DROP NOT NULL'))
        db.session.commit()
    except Exception:
        db.session.rollback()

    # Migração: mover base_data e groups_result do User para a Turma (se houver dados)
    try:
        users_with_data = User.query.filter((User.base_data != None) | (User.groups_result != None)).all()
        for u in users_with_data:
            # Se o usuário ainda não tem nenhuma turma, criamos a Turma Principal
            if not u.turmas:
                t = Turma(nome="Turma Principal", user_id=u.id, base_data=u.base_data, groups_result=u.groups_result)
                db.session.add(t)
            u.base_data = None
            u.groups_result = None
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        print(f"Erro na migração de Turmas: {e}")

# --- Funções de Autenticação e Decorators ---
def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'username' not in session:
            return redirect(url_for('login', next=request.url))
        return f(*args, **kwargs)
    return decorated_function

def get_user_data_folder():
    """Retorna o caminho para a pasta de dados do usuário logado (usado temporariamente para processar uploads)."""
    if 'username' in session:
        user_folder = os.path.join(app.config['SESSION_DATA_FOLDER'], session['username'])
        os.makedirs(user_folder, exist_ok=True)
        return user_folder
    return None

def grupos_para_dataframe(grupos):
    linhas = []
    for g in grupos:
        membros_nomes = [_get_nome_membro(m) for m in g.get("Membros", [])]
        for m in sorted(membros_nomes):
            linhas.append({"Grupo": g.get("Grupo"), "Dia": g.get("Dia"), "Horário": g.get("Horário"), "Aluno": m})
    return pd.DataFrame(linhas)

def montar_pdf_grupos(grupos):
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    _, altura = A4; margem = 2*cm; y = altura - margem
    c.setFont("Helvetica-Bold", 14)
    c.drawString(margem, y, "Grupos Formados")
    y -= 1.0 * cm
    for g in sorted(grupos, key=lambda x: x.get("Grupo", 0)):
        header = f"Grupo {g.get('Grupo')} — {g.get('Dia')} ({g.get('Horário')})"
        c.setFont("Helvetica-Bold", 11); c.drawString(margem, y, header); y -= 0.6 * cm
        c.setFont("Helvetica", 10)
        membros_nomes = [_get_nome_membro(m) for m in g.get("Membros", [])]
        for m in sorted(membros_nomes):
            c.drawString(margem + 0.5 * cm, y, f"• {m}"); y -= 0.5 * cm
            if y <= margem: c.showPage(); y = altura - margem
        y -= 0.3 * cm
        if y <= margem: c.showPage(); y = altura - margem
    c.save(); buf.seek(0)
    return buf

def _get_nome_membro(membro: any) -> str:
    if isinstance(membro, dict):
        return membro.get('nome', '')
    return str(membro)

# --- Auxiliares de E-mail e Tokens ---
def get_serializer():
    return URLSafeTimedSerializer(app.config['SECRET_KEY'])

def generate_reset_token(email):
    s = get_serializer()
    return s.dumps(email, salt='password-reset-salt')

def verify_reset_token(token, expiration=900): # 15 minutos
    s = get_serializer()
    try:
        email = s.loads(token, salt='password-reset-salt', max_age=expiration)
        return email
    except (SignatureExpired, BadSignature):
        return None

def enviar_email_recuperacao(email_destino, link_recuperacao):
    # Tentar via SendGrid Web API se SENDGRID_API_KEY estiver presente (evita bloqueios de porta SMTP do Render Free)
    sendgrid_api_key = os.environ.get('SENDGRID_API_KEY')
    if sendgrid_api_key:
        import requests
        url = "https://api.sendgrid.com/v3/mail/send"
        headers = {
            "Authorization": f"Bearer {sendgrid_api_key}",
            "Content-Type": "application/json"
        }
        mail_username = os.environ.get('MAIL_USERNAME', 'noreply.academysync@gmail.com')
        mail_sender_name = "AcademySync"
        
        payload = {
            "personalizations": [
                {
                    "to": [{"email": email_destino}],
                    "subject": "Recuperação de Senha - AcademySync"
                }
            ],
            "from": {
                "email": mail_username,
                "name": mail_sender_name
            },
            "content": [
                {
                    "type": "text/plain",
                    "value": f"Olá,\n\nVocê solicitou a redefinição de sua senha no AcademySync.\nPara redefinir sua senha, clique no link abaixo (válido por 15 minutos):\n\n{link_recuperacao}\n\nSe você não solicitou esta redefinição, por favor ignore este e-mail.\n\nAtenciosamente,\nEquipe AcademySync"
                }
            ]
        }
        try:
            res = requests.post(url, json=payload, headers=headers, timeout=10)
            if res.status_code in (200, 201, 202):
                print(f"E-mail via SendGrid enviado com sucesso para {email_destino}!")
                return True
            else:
                print(f"Erro SendGrid API (Status {res.status_code}): {res.text}")
        except Exception as e:
            print(f"Falha ao chamar SendGrid API: {e}")

    # Fallback para o SMTP clássico
    mail_server = os.environ.get('MAIL_SERVER')
    mail_port = os.environ.get('MAIL_PORT')
    mail_username = os.environ.get('MAIL_USERNAME')
    mail_password = os.environ.get('MAIL_PASSWORD')
    mail_sender = os.environ.get('MAIL_DEFAULT_SENDER', mail_username)
    
    if not all([mail_server, mail_port, mail_username, mail_password]):
        print("AVISO: Configurações de SMTP incompletas. Não foi possível enviar o e-mail de recuperação.")
        return False
        
    msg = EmailMessage()
    msg['Subject'] = 'Recuperação de Senha - AcademySync'
    msg['From'] = mail_sender
    msg['To'] = email_destino
    
    corpo = f"""Olá,

Você solicitou a redefinição de sua senha no AcademySync.
Para redefinir sua senha, clique no link abaixo (válido por 15 minutos):

{link_recuperacao}

Se você não solicitou esta redefinição, por favor ignore este e-mail.

Atenciosamente,
Equipe AcademySync
"""
    msg.set_content(corpo)
    
    try:
        port = int(mail_port)
        if port == 465:
            server = smtplib.SMTP_SSL(mail_server, port, timeout=10)
        else:
            server = smtplib.SMTP(mail_server, port, timeout=10)
            server.starttls()
            
        server.login(mail_username, mail_password)
        server.send_message(msg)
        server.quit()
        return True
    except Exception as e:
        print(f"ERRO ao enviar e-mail de recuperação para {email_destino}: {e}")
        return False

def _normalize_name_for_match(name: str) -> str:
    if not isinstance(name, str):
        return ''
    return name.replace('_', ' ').strip().lower()

def sincronizar_tabelas_dos_alunos(turma_id, base_df):
    if base_df is None or base_df.empty:
        return
    try:
        alunos_novos_nomes = set(base_df['aluno'].unique())
        
        # Obter alunos existentes no banco para esta turma
        alunos_existentes = Aluno.query.filter_by(turma_id=turma_id).all()
        mapeamento_existentes = {a.nome: a for a in alunos_existentes}
        
        # 1. Deletar alunos que não estão mais na base consolidada
        for nome_antigo, aluno_obj in list(mapeamento_existentes.items()):
            if nome_antigo not in alunos_novos_nomes:
                db.session.delete(aluno_obj)
                
        # 2. Inserir ou atualizar alunos novos validando conflitos
        for nome_novo in alunos_novos_nomes:
            # Obter matrícula do DataFrame, se a coluna existir
            matricula_val = None
            if 'matricula' in base_df.columns:
                aluno_rows = base_df[base_df['aluno'] == nome_novo]
                if not aluno_rows.empty:
                    val = str(aluno_rows.iloc[0]['matricula']).strip()
                    if val != '' and val.lower() != 'nan' and val.lower() != 'none':
                        matricula_val = val

            # Busca por correspondência no banco de dados
            match_nome = mapeamento_existentes.get(nome_novo)
            match_matricula = next((a for a in alunos_existentes if a.matricula == matricula_val), None) if matricula_val else None

            # REGRA 1: Mesmo Nome, Matrícula Diferente (Troca de matrícula)
            if match_nome and match_nome.matricula != matricula_val:
                flash(
                    f"Aviso: O aluno '{nome_novo.replace('_', ' ')}' já está cadastrado, "
                    f"mas a matrícula no PDF ({matricula_val}) diverge do banco ({match_nome.matricula}). "
                    f"A matrícula foi atualizada.", "warning"
                )
                match_nome.matricula = matricula_val

            # REGRA 2: Matrículas Iguais, Nomes Diferentes (Troca de nome)
            elif match_matricula and match_matricula.nome.lower() != nome_novo.lower():
                nome_antigo_limpo = match_matricula.nome.replace('_', ' ')
                flash(
                    f"Aviso: A matrícula {matricula_val} já pertence a '{nome_antigo_limpo}', "
                    f"mas o PDF indicou o nome '{nome_novo.replace('_', ' ')}'. "
                    f"O nome do aluno foi atualizado.", "warning"
                )
                match_matricula.nome = nome_novo
                match_matricula.nome_custom = nome_novo.replace('_', ' ')

            # REGRA 3: Nome e Matrícula Diferentes (Adiciona Novo Aluno)
            elif not match_nome and not match_matricula:
                novo_aluno = Aluno(
                    nome=nome_novo,
                    nome_custom=nome_novo.replace('_', ' '),
                    turma_id=turma_id,
                    matricula=matricula_val
                )
                db.session.add(novo_aluno)
                
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        print(f"Erro ao sincronizar alunos: {e}")

def sincronizar_tabelas_dos_grupos(turma_id, result):
    try:
        # a. Resetar todos os grupo_id dos alunos dessa turma para None
        Aluno.query.filter_by(turma_id=turma_id).update({Aluno.grupo_id: None})
        
        # b. Deletar todos os grupos antigos dessa turma
        Grupo.query.filter_by(turma_id=turma_id).delete()
        db.session.commit()
        
        if not result or not result.get('grupos'):
            return
            
        # c. Criar os novos grupos e associar os alunos
        for g in result.get('grupos', []):
            g_nome = f"Grupo {g['Grupo']}" if isinstance(g.get('Grupo'), (int, float)) else str(g.get('Grupo'))
            grupo_db = Grupo(
                turma_id=turma_id,
                nome=g_nome,
                dia=g.get('Dia', ''),
                horario=g.get('Horário', '')
            )
            db.session.add(grupo_db)
            db.session.flush() # para obter o grupo_db.id
            
            for membro_item in g.get('Membros', []):
                nome_membro = _get_nome_membro(membro_item)
                # Encontrar aluno e associar
                aluno_db = Aluno.query.filter_by(turma_id=turma_id, nome=nome_membro).first()
                if not aluno_db:
                    # Tentar busca normalizada
                    all_alunos = Aluno.query.filter_by(turma_id=turma_id).all()
                    for al in all_alunos:
                        if _normalize_name_for_match(al.nome) == _normalize_name_for_match(nome_membro):
                            aluno_db = al
                            break
                if aluno_db:
                    aluno_db.grupo_id = grupo_db.id
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        print(f"Erro ao sincronizar grupos: {e}")

# --- Rotas de Autenticação ---
@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form['username']
        email = request.form['email']
        password = request.form['password']
        confirm_password = request.form.get('confirm_password')
        
        if password != confirm_password:
            flash('As senhas não coincidem.', 'danger')
            return redirect(url_for('register'))
        
        # Obter pergunta de segurança
        security_q = request.form.get('security_question')
        if security_q == 'custom':
            security_q = request.form.get('custom_question_text')
        
        security_ans = request.form.get('security_answer')

        if not security_q or not security_ans:
            flash('Pergunta e resposta de segurança são obrigatórias.', 'danger')
            return redirect(url_for('register'))

        user_exists = User.query.filter_by(username=username).first()
        if user_exists:
            flash('Este nome de usuário já existe.', 'danger')
            return redirect(url_for('register'))

        email_exists = User.query.filter_by(email=email).first()
        if email_exists:
            flash('Este e-mail já está cadastrado.', 'danger')
            return redirect(url_for('register'))

        hashed_password = generate_password_hash(password)
        hashed_answer = generate_password_hash(security_ans.strip().lower())
        
        new_user = User(
            username=username, 
            email=email, 
            password_hash=hashed_password,
            security_question=security_q,
            security_answer_hash=hashed_answer
        )
        db.session.add(new_user)
        db.session.commit()

        flash('Conta criada com sucesso! Faça o login.', 'success')
        return redirect(url_for('login'))
    return render_template('register.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username_or_email = request.form['username_or_email'].strip()
        password = request.form['password']
        
        # Buscar por username ou email
        user = User.query.filter((User.username == username_or_email) | (User.email == username_or_email)).first()

        if user and check_password_hash(user.password_hash, password):
            session['username'] = user.username # Mantém username na sessão para compatibilidade
            flash('Login bem-sucedido!', 'success')
            return redirect(url_for('index'))
        else:
            flash('Usuário ou senha inválidos.', 'error')
            
    return render_template('login.html')

@app.route('/change_password', methods=['POST'])
@login_required
def change_password():
    username = session.get('username')
    current_password = request.form.get('current_password')
    new_password = request.form.get('new_password')
    confirm_password = request.form.get('confirm_password')

    if new_password != confirm_password:
        flash('A nova senha e a confirmação não correspondem.', 'danger')
        return redirect(url_for('index'))

    user = User.query.filter_by(username=username).first()

    if user and check_password_hash(user.password_hash, current_password):
        user.password_hash = generate_password_hash(new_password)
        db.session.commit()
        flash('Senha alterada com sucesso!', 'success')
    else:
        flash('A senha atual está incorreta.', 'danger')

    return redirect(url_for('index'))

@app.route('/configure_security', methods=['POST'])
@login_required
def configure_security():
    username = session.get('username')
    security_q = request.form.get('security_question')
    if security_q == 'custom':
        security_q = request.form.get('custom_question_text')
        
    security_ans = request.form.get('security_answer')
    confirm_pwd = request.form.get('confirm_password_security')
    
    if not security_q or not security_ans or not confirm_pwd:
        flash('Todos os campos são obrigatórios.', 'danger')
        return redirect(url_for('index'))
        
    user = User.query.filter_by(username=username).first()
    
    if user and check_password_hash(user.password_hash, confirm_pwd):
        user.security_question = security_q
        user.security_answer_hash = generate_password_hash(security_ans.strip().lower())
        db.session.commit()
        flash('Pergunta de segurança configurada com sucesso!', 'success')
    else:
        flash('Senha de confirmação incorreta.', 'danger')
        
    return redirect(url_for('index'))

@app.route('/setup_recovery_question', methods=['POST'])
@login_required
def setup_recovery_question():
    username = session.get('username')
    security_q = request.form.get('security_question')
    if security_q == 'custom':
        security_q = request.form.get('custom_question_text')
        
    security_ans = request.form.get('security_answer')
    
    if not security_q or not security_ans:
        flash('Pergunta e resposta de segurança são obrigatórias.', 'danger')
        return redirect(url_for('index'))
        
    user = User.query.filter_by(username=username).first()
    if user:
        user.security_question = security_q
        user.security_answer_hash = generate_password_hash(security_ans.strip().lower())
        db.session.commit()
        flash('Pergunta de segurança configurada com sucesso!', 'success')
    else:
        flash('Erro ao processar a solicitação.', 'danger')
        
    return redirect(url_for('index'))

@app.route('/logout')
def logout():
    session.clear()
    flash('Você foi desconectado.', 'info')
    return redirect(url_for('login'))

@app.route('/forgot_password', methods=['GET', 'POST'])
def forgot_password():
    step = 1
    username_or_email = ""
    question = ""
    
    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'identify':
            username_or_email = request.form.get('username_or_email', '').strip()
            # Procurar por username ou email
            user = User.query.filter((User.username == username_or_email) | (User.email == username_or_email)).first()
            
            if user:
                if not user.security_question or not user.security_answer_hash:
                    # Verificar se estamos rodando em ambiente local de desenvolvimento para gerar link de bypass
                    host_local = request.host.split(':')[0] in ('127.0.0.1', 'localhost')
                    is_local = app.config['SQLALCHEMY_DATABASE_URI'].startswith('sqlite') or host_local
                    
                    if is_local:
                        token = generate_reset_token(user.email)
                        link_recuperacao = url_for('reset_password', token=token, _external=True)
                        flash('Sua conta não possui pergunta de segurança configurada. (Modo de Demonstração Local: use o link abaixo).', 'warning')
                        return render_template('forgot_password.html', link_recuperacao=link_recuperacao, local_debug=True, step=1)
                    else:
                        flash('Esta conta foi criada antes do recurso de perguntas de segurança. Por favor, contate o administrador para redefinir sua senha.', 'danger')
                        return render_template('forgot_password.html', step=1)
                else:
                    # Ir para o passo 2
                    return render_template('forgot_password.html', step=2, username=user.username, question=user.security_question)
            else:
                flash('Usuário ou e-mail não cadastrado.', 'danger')
                return render_template('forgot_password.html', step=1)
                
        elif action == 'verify':
            username = request.form.get('username', '').strip()
            answer = request.form.get('security_answer', '').strip()
            
            user = User.query.filter_by(username=username).first()
            if user:
                if user.security_answer_hash and check_password_hash(user.security_answer_hash, answer.lower()):
                    # Resposta correta! Gera o token de reset e redireciona
                    token = generate_reset_token(user.email)
                    return redirect(url_for('reset_password', token=token))
                else:
                    flash('Resposta incorreta. Tente novamente.', 'danger')
                    return render_template('forgot_password.html', step=2, username=user.username, question=user.security_question)
            else:
                flash('Erro ao processar a solicitação.', 'danger')
                return redirect(url_for('forgot_password'))
                
    return render_template('forgot_password.html', step=step)

@app.route('/debug_smtp')
def debug_smtp():
    import smtplib
    from email.message import EmailMessage
    import traceback
    
    mail_server = os.environ.get('MAIL_SERVER')
    mail_port = os.environ.get('MAIL_PORT')
    mail_username = os.environ.get('MAIL_USERNAME')
    mail_password = os.environ.get('MAIL_PASSWORD')
    mail_sender = os.environ.get('MAIL_DEFAULT_SENDER', mail_username)
    
    smtp_test_result = "Not tested"
    
    if all([mail_server, mail_port, mail_username, mail_password]):
        try:
            msg = EmailMessage()
            msg['Subject'] = 'Diagnostic SMTP - AcademySync'
            msg['From'] = mail_sender
            msg['To'] = mail_username
            msg.set_content("Teste de diagnóstico SMTP no Render.")
            
            port = int(mail_port)
            if port == 465:
                server = smtplib.SMTP_SSL(mail_server, port, timeout=10)
            else:
                server = smtplib.SMTP(mail_server, port, timeout=10)
                server.starttls()
                
            server.login(mail_username, mail_password)
            server.send_message(msg)
            server.quit()
            smtp_test_result = "Success!"
        except Exception as e:
            smtp_test_result = f"Failed: {str(e)} - {traceback.format_exc()}"
            
    return json.dumps({
        "MAIL_SERVER": mail_server,
        "MAIL_PORT": mail_port,
        "MAIL_USERNAME": mail_username,
        "HAS_PASSWORD": bool(mail_password),
        "PASSWORD_LENGTH": len(mail_password) if mail_password else 0,
        "SMTP_TEST_RESULT": smtp_test_result
    }), 200

@app.route('/reset_password/<token>', methods=['GET', 'POST'])
def reset_password(token):
    email = verify_reset_token(token)
    if not email:
        flash('O link de recuperação é inválido ou expirou.', 'danger')
        return redirect(url_for('forgot_password'))
        
    if request.method == 'POST':
        new_password = request.form['new_password']
        confirm_password = request.form['confirm_password']
        
        if new_password != confirm_password:
            flash('As senhas não correspondem.', 'danger')
            return render_template('reset_password.html', token=token)
            
        user = User.query.filter_by(email=email).first()
        if user:
            user.password_hash = generate_password_hash(new_password)
            db.session.commit()
            flash('Sua senha foi redefinida com sucesso!', 'success')
            return redirect(url_for('login'))
        else:
            flash('Usuário não encontrado.', 'danger')
            return redirect(url_for('login'))
            
    return render_template('reset_password.html', token=token)


# --- API de Horários para UI/UX Dinâmica ---

@app.route('/aluno/schedules_json/<int:turma_id>')
@login_required
def get_schedules_json(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.base_data:
        return json.dumps({}), 200
        
    try:
        df = pd.read_csv(BytesIO(turma.base_data))
        
        # Mapear dias da semana
        dias_completos = {'Seg': 'Segunda', 'Ter': 'Terça', 'Qua': 'Quarta', 'Qui': 'Quinta', 'Sex': 'Sexta', 'Sab': 'Sábado'}
        
        # Celulas ocupadas por aluno
        schedules = {}
        for aluno in df['aluno'].unique():
            schedules[str(aluno).replace('_', ' ')] = {}

        for _, row in df.iterrows():
            materia = row['materia']
            if pd.isna(materia) or str(materia).strip() in ('', '---'):
                continue

            intervalo = obter_intervalo_minutos(row['horario'])
            if intervalo is None:
                continue

            nome_limpo = str(row['aluno']).replace('_', ' ')
            dia_completo = dias_completos.get(row['dia'], row['dia'])
            materia_str = str(materia).strip()
            
            # Mapear status das siglas/descrições
            if materia_str in ('Acompanhamento Comportamental', 'AC') or 'Acompanhamento Comportamental' in materia_str:
                status = 'ac'
            elif materia_str in ('Psicoterapia', 'PSI') or 'Psicoterapia' in materia_str:
                status = 'psi'
            elif materia_str in ('Imersão', 'IM') or 'Imersão' in materia_str:
                status = 'im'
            elif materia_str in ('Acompanhamento Pedagógico', 'AP') or 'Acompanhamento Pedagógico' in materia_str:
                status = 'ap'
            elif materia_str in ('Realocado', 'R') or 'Realocado' in materia_str:
                status = 'r'
            elif materia_str.startswith('Custom:'):
                status = materia_str[7:].strip()
            else:
                # Qualquer matéria trazida de PDF (Cálculo, Estatística, Aula de Graduação, X, etc.) é Aula de Graduação
                status = 'class'

            ini, fim = intervalo

            aluno_sched = schedules.get(nome_limpo)
            if aluno_sched is None:
                continue

            for slot in slots_ocupados_no_intervalo(ini, fim):
                aluno_sched[f"{dia_completo}_{slot}"] = status

        return json.dumps(schedules), 200
    except Exception as e:
        return json.dumps({"error": str(e)}), 500


@app.route('/aluno/all_schedules_json')
@login_required
def get_all_schedules_json():
    user = User.query.filter_by(username=session['username']).first()
    if not user:
        return json.dumps({}), 200
        
    turmas = Turma.query.filter_by(user_id=user.id).all()
    schedules = {}
    student_turma_map = {}
    student_turma_id_map = {}
    
    dias_completos = {'Seg': 'Segunda', 'Ter': 'Terça', 'Qua': 'Quarta', 'Qui': 'Quinta', 'Sex': 'Sexta', 'Sab': 'Sábado'}
    
    for t in turmas:
        if not t.base_data:
            continue
        try:
            df = pd.read_csv(BytesIO(t.base_data))
            
            # Inicializar alunos desta turma
            alunos_unicos = sorted(df['aluno'].unique())
            for aluno in alunos_unicos:
                nome_limpo = aluno.replace('_', ' ')
                schedules[nome_limpo] = {}
                student_turma_map[nome_limpo] = t.nome
                student_turma_id_map[nome_limpo] = t.id
                # Inicializar com free para todos os slots de 30 minutos
                for dia in ['Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado']:
                    for slot in HORARIOS_GRADE:
                        schedules[nome_limpo][f"{dia}_{slot}"] = 'free'
                        
            # Preencher do DataFrame
            for _, row in df.iterrows():
                materia = row['materia']
                if pd.isna(materia) or str(materia).strip() in ('', '---'):
                    continue

                intervalo = obter_intervalo_minutos(row['horario'])
                if intervalo is None:
                    continue

                nome_limpo = str(row['aluno']).replace('_', ' ')
                dia_completo = dias_completos.get(row['dia'], row['dia'])
                materia_str = str(materia).strip()
                
                # Mapear status das siglas/descrições
                if materia_str in ('Acompanhamento Comportamental', 'AC') or 'Acompanhamento Comportamental' in materia_str:
                    status = 'ac'
                elif materia_str in ('Psicoterapia', 'PSI') or 'Psicoterapia' in materia_str:
                    status = 'psi'
                elif materia_str in ('Imersão', 'IM') or 'Imersão' in materia_str:
                    status = 'im'
                elif materia_str in ('Acompanhamento Pedagógico', 'AP') or 'Acompanhamento Pedagógico' in materia_str:
                    status = 'ap'
                elif materia_str in ('Realocado', 'R') or 'Realocado' in materia_str:
                    status = 'r'
                elif materia_str.startswith('Custom:'):
                    status = materia_str[7:].strip()
                else:
                    # Qualquer matéria trazida de PDF (Cálculo, Estatística, Aula de Graduação, X, etc.) é Aula de Graduação
                    status = 'class'
                ini, fim = intervalo

                aluno_sched = schedules.get(nome_limpo)
                if aluno_sched is None:
                    continue

                for slot in slots_ocupados_no_intervalo(ini, fim):
                    aluno_sched[f"{dia_completo}_{slot}"] = status
        except Exception as e:
            print(f"Erro ao processar turma {t.id} na rota consolidada: {e}")
            
    lista_alunos = sorted(list(schedules.keys()))
    payload = {
        "schedules": schedules,
        "student_turma_map": student_turma_map,
        "student_turma_id_map": student_turma_id_map,
        "lista_alunos": lista_alunos
    }
    return json.dumps(payload), 200

@app.route('/aluno/update_schedule/<int:turma_id>', methods=['POST'])
@login_required
def update_student_schedule(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.base_data:
        return json.dumps({"error": "Dados não encontrados"}), 404
        
    data = request.json
    aluno_nome = data.get('aluno')
    dia = data.get('dia')
    horario = str(data.get('horario', '')).strip()  # ex.: "07:30"
    novo_status = data.get('status')
    force_swap = data.get('force_swap', False)

    if horario not in HORARIOS_GRADE:
        return json.dumps({"error": "Horário inválido"}), 400

    dia_abrev_map = {'Segunda': 'Seg', 'Terça': 'Ter', 'Quarta': 'Qua', 'Quinta': 'Qui', 'Sexta': 'Sex', 'Sábado': 'Sab'}
    dia_abrev = dia_abrev_map.get(dia, dia)
    dias_completos = {'Seg': 'Segunda', 'Ter': 'Terça', 'Qua': 'Quarta', 'Qui': 'Quinta', 'Sex': 'Sexta', 'Sab': 'Sábado'}

    # Validação de conflitos para compromissos restritos
    RESTRICTED_MAP = {
        'ac': 'Acompanhamento Comportamental',
        'psi': 'Psicoterapia',
        'ap': 'Acompanhamento Pedagógico',
        'r': 'Realocado'
    }

    conflito_aluno = None
    conflito_turma_id = None
    conflito_turma_name = None

    if novo_status in RESTRICTED_MAP:
        turmas = Turma.query.filter_by(user_id=user.id).all()
        
        def slot_matches(h, target_horario, target_horario_salvar):
            if str(h).strip() == target_horario_salvar:
                return True
            interval = obter_intervalo_minutos(h)
            if interval:
                ini, fim = interval
                return target_horario in slots_ocupados_no_intervalo(ini, fim)
            return False

        horario_salvar_temp = format_horario_slot(horario)

        for t in turmas:
            if not t.base_data:
                continue
            try:
                other_df = pd.read_csv(BytesIO(t.base_data))
                for _, row in other_df.iterrows():
                    materia = row['materia']
                    if pd.isna(materia):
                        continue
                    
                    materia_str = str(materia).strip()
                    line_status = None
                    if materia_str in ('Acompanhamento Comportamental', 'AC') or 'Acompanhamento Comportamental' in materia_str:
                        line_status = 'ac'
                    elif materia_str in ('Psicoterapia', 'PSI') or 'Psicoterapia' in materia_str:
                        line_status = 'psi'
                    elif materia_str in ('Acompanhamento Pedagógico', 'AP') or 'Acompanhamento Pedagógico' in materia_str:
                        line_status = 'ap'
                    elif materia_str in ('Realocado', 'R') or 'Realocado' in materia_str:
                        line_status = 'r'
                        
                    if line_status == novo_status:
                        line_aluno_limpo = str(row['aluno']).replace('_', ' ')
                        if line_aluno_limpo.lower() != aluno_nome.lower():
                            if dias_completos.get(row['dia'], row['dia']) == dia:
                                if slot_matches(row['horario'], horario, horario_salvar_temp):
                                    conflito_aluno = line_aluno_limpo
                                    conflito_turma_id = t.id
                                    conflito_turma_name = t.nome
                                    break
                if conflito_aluno:
                    break
            except Exception as e:
                print(f"Erro ao verificar conflito: {e}")

        if conflito_aluno:
            if not force_swap:
                return json.dumps({
                    "conflict": True,
                    "conflicting_student": conflito_aluno,
                    "conflicting_turma_name": conflito_turma_name,
                    "conflicting_turma_id": conflito_turma_id,
                    "materia_nome": RESTRICTED_MAP[novo_status]
                }), 409
            else:
                try:
                    aluno_conflito_key = conflito_aluno.replace(' ', '_')
                    horario_salvar_temp = format_horario_slot(horario)
                    
                    if conflito_turma_id != turma_id:
                        conf_turma = Turma.query.get(conflito_turma_id)
                        conf_df = pd.read_csv(BytesIO(conf_turma.base_data))
                        
                        cond_horario_conf = conf_df['horario'].apply(lambda h: slot_matches(h, horario, horario_salvar_temp))
                        cond_conf = (
                            conf_df['aluno'].apply(lambda a: str(a).replace(' ', '_').strip().lower() == aluno_conflito_key.lower())
                            & (conf_df['dia'] == dia_abrev)
                            & cond_horario_conf
                        )
                        if not conf_df[cond_conf].empty:
                            conf_df.loc[cond_conf, 'materia'] = '---'
                            
                        out_conf = BytesIO()
                        conf_df.to_csv(out_conf, index=False, encoding='utf-8')
                        conf_turma.base_data = out_conf.getvalue()
                except Exception as e:
                    print(f"Erro ao realizar swap em outra turma: {e}")
    
    try:
        df = pd.read_csv(BytesIO(turma.base_data))
        aluno_key = aluno_nome.replace(' ', '_')
        # Match "João Silva" e "João_Silva" como o mesmo aluno
        def _aluno_keys_equal(a, target_key):
            return str(a).replace(' ', '_').strip().lower() == target_key.lower()

        horario_salvar = format_horario_slot(horario)

        if force_swap and conflito_aluno and conflito_turma_id == turma_id:
            aluno_conflito_key = conflito_aluno.replace(' ', '_')
            cond_horario_conf = df['horario'].apply(lambda h: str(h).strip() == horario_salvar or (obter_intervalo_minutos(h) and horario in slots_ocupados_no_intervalo(obter_intervalo_minutos(h)[0], obter_intervalo_minutos(h)[1])))
            cond_conf = df['aluno'].apply(lambda a: _aluno_keys_equal(a, aluno_conflito_key)) & (df['dia'] == dia_abrev) & cond_horario_conf
            if not df[cond_conf].empty:
                df.loc[cond_conf, 'materia'] = '---'

        def _match_horario(h):
            if str(h).strip() == horario_salvar:
                return True
            return obter_hora_inicio(h) == horario

        cond_horario = df['horario'].apply(_match_horario)
        cond_aluno = df['aluno'].apply(lambda a: _aluno_keys_equal(a, aluno_key))
        # Normaliza chave no CSV (espaços → '_') para não duplicar na UI
        df.loc[cond_aluno, 'aluno'] = aluno_key
        condicao = cond_aluno & (df['dia'] == dia_abrev) & cond_horario
        
        materia_val = '---'
        if novo_status == 'class':
            materia_val = 'Aula de Graduação'
        elif novo_status == 'ac':
            materia_val = 'Acompanhamento Comportamental'
        elif novo_status == 'psi':
            materia_val = 'Psicoterapia'
        elif novo_status == 'im':
            materia_val = 'Imersão'
        elif novo_status == 'ap':
            materia_val = 'Acompanhamento Pedagógico'
        elif novo_status == 'r':
            materia_val = 'Realocado'
        elif novo_status != 'free':
            materia_val = f"Custom: {novo_status}" if not novo_status.startswith('Custom:') else novo_status
            
        if not df[condicao].empty:
            df.loc[condicao, 'materia'] = materia_val
        else:
            new_row = pd.DataFrame([{
                'aluno': aluno_key,
                'dia': dia_abrev,
                'horario': horario_salvar,
                'materia': materia_val
            }])
            df = pd.concat([df, new_row], ignore_index=True)
            
        # Salvar CSV atualizado
        csv_buffer = BytesIO()
        df.to_csv(csv_buffer, index=False, encoding='utf-8')

        turma.base_data = csv_buffer.getvalue()
        
        # Se houver grupos formados, atualizar o cache de compatibilidade e detectar conflito
        conflito_gerado = False
        grupo_afetado = None
        
        if turma.groups_result:
            result = pickle.loads(turma.groups_result)
            slots_por_aluno = result.get('slots_por_aluno', {})
            
            # Recomputar a disponibilidade deste aluno para todos os blocos cadastrados
            slots_por_aluno[aluno_key] = set()
            blocos = result.get('blocos', [])
            
            # Obter as aulas ocupadas do aluno do DataFrame atualizado
            aluno_df = df[df['aluno'].str.lower() == aluno_key.lower()]
            aluno_busy = aluno_df[aluno_df['materia'] != '---']
            
            # Agrupar horários ocupados por dia
            ocupados_dia = {}
            for _, row in aluno_busy.iterrows():
                d_val = str(row['dia'])
                if d_val not in ocupados_dia:
                    ocupados_dia[d_val] = []
                ocupados_dia[d_val].append(row['horario'])
                
            for b in blocos:
                nome_bloco = b.get("nome")
                ini = int(b["inicio"])
                fim = int(b["fim"])
                dias = b.get("dias") or b.get("dia") or []
                if isinstance(dias, str):
                    dias = [d.strip() for d in dias.split(",") if d.strip()]
                label = f"{nome_bloco} ({ini:02d}h-{fim:02d}h)"
                
                for dia in dias:
                    busy_hours = ocupados_dia.get(dia, [])
                    livre = True
                    for h in busy_hours:
                        if horarios_conflitam(h, ini, fim):
                            livre = False
                            break
                    if livre:
                        slots_por_aluno[aluno_key].add((dia, label))
            
            # Encontrar grupo do aluno
            grupos = result.get('grupos', [])
            aluno_grupo = None
            for g in grupos:
                membros_nomes = [_get_nome_membro(m) for m in g.get('Membros', [])]
                if aluno_nome.lower() in [m.lower() for m in membros_nomes]:
                    aluno_grupo = g
                    break
            
            if aluno_grupo and novo_status != 'free':
                g_dia = aluno_grupo.get('Dia')
                g_horario_str = aluno_grupo.get('Horário')
                if g_dia == dia_abrev:
                    try:
                        # Extrai horas usando regex (compatível com os blocos)
                        match = re.search(r'(\d+)h-(\d+)h', g_horario_str)
                        if match:
                            h_inicio = int(match.group(1))
                            h_fim = int(match.group(2))
                        else:
                            h_inicio = int(g_horario_str.split(' - ')[0].split(':')[0])
                            h_fim = int(g_horario_str.split(' - ')[1].split(':')[0])
                        if horario >= h_inicio and horario < h_fim:
                            conflito_gerado = True
                            grupo_afetado = aluno_grupo.get('Grupo')
                    except Exception:
                        pass
            
            result['slots_por_aluno'] = slots_por_aluno
            turma.groups_result = pickle.dumps(result)
            
            # Sincronizar tabelas relacionais de grupos e alunos
            sincronizar_tabelas_dos_grupos(turma.id, result)
            
        db.session.commit()
        return json.dumps({
            "success": True,
            "conflito": conflito_gerado,
            "grupo": grupo_afetado
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return json.dumps({"error": str(e)}), 500


# --- Rotas Principais ---
@app.route('/', methods=['GET', 'POST'])
@login_required # Este decorator FORÇA o login para acessar a esta página
def index():
    user = User.query.filter_by(username=session['username']).first()
    if not user:
        return redirect(url_for('login'))
        
    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'create_turma':
            nome = request.form.get('turma_nome', '').strip()
            if not nome:
                flash('O nome da turma não pode ser vazio.', 'error')
                return redirect(url_for('index'))
                
            em_imersao = (request.form.get('em_imersao') == 'true')
            new_turma = Turma(nome=nome, user_id=user.id, em_imersao=em_imersao)
            db.session.add(new_turma)
            db.session.commit() # Salva primeiro para gerar o ID
            
            # Processar PDFs se houver
            files = request.files.getlist('pdf_files')
            if files and files[0].filename != '':
                filepaths = []
                user_data_folder = get_user_data_folder()
                upload_folder_user = os.path.join(user_data_folder, 'uploads') if user_data_folder else 'uploads'
                os.makedirs(upload_folder_user, exist_ok=True)

                for file in files:
                    filename = secure_filename(file.filename)
                    filepath = os.path.join(upload_folder_user, filename)
                    file.save(filepath)
                    filepaths.append(filepath)
                
                base_df = processar_pdfs_para_dataframe(filepaths)
                
                # Limpar PDFs temporários
                try:
                    for fpath in filepaths:
                        if os.path.exists(fpath):
                            os.remove(fpath)
                except Exception as e:
                    print(f"AVISO: Falha ao limpar PDFs temporários: {e}")
                
                if base_df is not None and not base_df.empty:
                    csv_buffer = BytesIO()
                    base_df.to_csv(csv_buffer, index=False, encoding='utf-8')
                    new_turma.base_data = csv_buffer.getvalue()
                    db.session.commit()
                    
                    save_studant(base_df, new_turma.id)
                    flash(f'Turma "{nome}" criada com {len(base_df["aluno"].unique())} alunos do(s) PDF(s) com sucesso!', 'success')
                else:
                    flash(f'Turma "{nome}" criada vazia. Não foi possível extrair dados dos PDFs.', 'warning')
            else:
                flash(f'Turma "{nome}" criada com sucesso!', 'success')
                
            return redirect(url_for('index'))
            
        elif action == 'delete_turma':
            turma_id = request.form.get('turma_id')
            if not turma_id:
                flash('ID da turma inválido.', 'error')
                return redirect(url_for('index'))
                
            turma = Turma.query.filter_by(id=int(turma_id), user_id=user.id).first()
            if turma:
                db.session.delete(turma)
                db.session.commit()
                flash('Turma excluída com sucesso.', 'success')
            else:
                flash('Turma não encontrada.', 'error')
            return redirect(url_for('index'))

        elif action == 'edit_turma':
            turma_id = request.form.get('turma_id')
            nome = request.form.get('turma_nome', '').strip()
            if not turma_id or not nome:
                flash('ID da turma ou nome inválido.', 'error')
                return redirect(url_for('index'))
                
            turma = Turma.query.filter_by(id=int(turma_id), user_id=user.id).first()
            if turma:
                old_nome = turma.nome
                turma.nome = nome
                db.session.commit()
                flash(f'Turma "{old_nome}" alterada para "{nome}" com sucesso.', 'success')
            else:
                flash('Turma não encontrada.', 'error')
            return redirect(url_for('index'))

    # LÓGICA GET: Listar turmas
    turmas = Turma.query.filter_by(user_id=user.id).all()
    turmas_info = []
    for t in turmas:
        alunos_count = 0
        if t.base_data:
            try:
                df = pd.read_csv(BytesIO(t.base_data))
                alunos_count = len(df['aluno'].unique())
            except Exception:
                pass
        
        grupos_count = 0
        if t.groups_result:
            try:
                res = pickle.loads(t.groups_result)
                grupos_count = len(res.get('grupos', []))
            except Exception:
                pass
                
        turmas_info.append({
            "id": t.id,
            "nome": t.nome,
            "alunos_count": alunos_count,
            "grupos_count": grupos_count,
            "em_imersao": t.em_imersao
        })

    has_security_question = bool(user.security_question and user.security_answer_hash)
    return render_template('landing.html', 
                           turmas=turmas_info, 
                           has_security_question=has_security_question)

@app.route('/turma/<int:turma_id>', methods=['GET', 'POST'])
@login_required
def view_turma(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    if not user:
        return redirect(url_for('login'))
        
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma:
        flash('Turma não encontrada.', 'error')
        return redirect(url_for('index'))
        
    user_data_folder = get_user_data_folder()
    
    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'process_pdfs':
            files = request.files.getlist('pdf_files')
            if not files or files[0].filename == '':
                flash('Nenhum arquivo selecionado.', 'error')
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return jsonify({"success": True, "redirect_url": url_for('view_turma', turma_id=turma_id)})
                return redirect(url_for('view_turma', turma_id=turma_id))

            filepaths = []
            upload_folder_user = os.path.join(user_data_folder, 'uploads') if user_data_folder else 'uploads'
            os.makedirs(upload_folder_user, exist_ok=True)

            for file in files:
                filename = secure_filename(file.filename)
                filepath = os.path.join(upload_folder_user, filename)
                file.save(filepath)
                filepaths.append(filepath)
            
            base_df = processar_pdfs_para_dataframe(filepaths)
            
            # Limpa os PDFs temporários salvos no disco
            try:
                for fpath in filepaths:
                    if os.path.exists(fpath):
                        os.remove(fpath)
            except Exception as e:
                print(f"AVISO: Falha ao limpar PDFs temporários: {e}")
    
            if base_df is not None and not base_df.empty:
                # 1. Carrega dados existentes da turma, se houver, e realiza o merge cumulativo
                if turma.base_data:
                    try:
                        existing_df = pd.read_csv(BytesIO(turma.base_data))
                        
                        # Lista de alunos que estão sendo carregados nesta carga de PDFs
                        novos_alunos = base_df['aluno'].unique()
                        
                        # Remove os dados antigos desses alunos específicos para evitar duplicados
                        existing_df = existing_df[~existing_df['aluno'].isin(novos_alunos)]
                        
                        # Une os dados antigos preservados com os novos
                        base_df = pd.concat([existing_df, base_df], ignore_index=True)
                    except Exception as e:
                        print(f"Aviso ao ler base_data anterior para merge: {e}")

                # 2. Salva o DataFrame consolidado de volta no Banco de Dados
                csv_buffer = BytesIO()
                base_df.to_csv(csv_buffer, index=False, encoding='utf-8')
                turma.base_data = csv_buffer.getvalue()
                db.session.commit()
                
                # 3. Sincroniza alunos na tabela Aluno
                sincronizar_tabelas_dos_alunos(turma.id, base_df)
                
                flash(f'Dados processados com sucesso! Total na turma: {len(base_df["aluno"].unique())} aluno(s).', 'success')
            else:
                flash('Não foi possível extrair dados válidos dos PDFs.', 'error')
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return jsonify({"success": True, "redirect_url": url_for('view_turma', turma_id=turma_id)})
            return redirect(url_for('view_turma', turma_id=turma_id))

        elif action == 'generate_groups':
            if not turma.base_data:
                flash('Primeiro processe os arquivos PDF.', 'error')
                return redirect(url_for('view_turma', turma_id=turma_id))
            
            try:
                base_df = pd.read_csv(BytesIO(turma.base_data))
            except Exception as e:
                flash('Erro ao ler os dados de alunos salvos no banco. Processe os PDFs novamente.', 'error')
                turma.base_data = None
                db.session.commit()
                return redirect(url_for('view_turma', turma_id=turma_id))

            # Leitura dos blocos do formulário
            nomes_blocos = request.form.getlist('bloco_nome[]')
            horas_inicio = request.form.getlist('bloco_inicio[]')
            horas_fim = request.form.getlist('bloco_fim[]')
            dias_semana = request.form.getlist('bloco_dias[]')

            blocos = []
            if not all([nomes_blocos, horas_inicio, horas_fim, dias_semana]):
                flash('Erro ao ler os grupos. Tente novamente.', 'error')
                return redirect(url_for('view_turma', turma_id=turma_id))

            for i in range(len(nomes_blocos)):
                try:
                    dias = [dias_semana[i]] if dias_semana[i] else []
                    
                    bloco = {
                        "nome": nomes_blocos[i],
                        "inicio": int(horas_inicio[i]),
                        "fim": int(horas_fim[i]),
                        "dias": dias
                    }
                    if all(bloco.values()): # Garante que nenhum campo está vazio
                        blocos.append(bloco)
                except (ValueError, IndexError):
                    flash(f'Grupo inválido encontrado. Verifique os dados.', 'warning')
                    continue
            
            if not blocos:
                flash('Pelo menos um grupo válido é necessário.', 'error')
                return redirect(url_for('view_turma', turma_id=turma_id))

            num_grupos = len(blocos)
            result = formar_grupos(base_df, num_grupos=num_grupos, blocos=blocos)
            result['blocos'] = blocos
            
            alunos_sem_grupo = result.get('alunos_sem_grupo')
            if alunos_sem_grupo:
                display_limit = 5
                nomes_display = ", ".join(al.replace('_', ' ') for al in alunos_sem_grupo[:display_limit])
                extra_count = len(alunos_sem_grupo) - display_limit
                msg = (f"{len(alunos_sem_grupo)} aluno(s) não foram alocados por conflito de horário: {nomes_display}")
                if extra_count > 0:
                    msg += f" e mais {extra_count}."
                flash(msg, 'warning')

            # Salva o resultado dos grupos no Banco de Dados em formato Pickle binário
            turma.groups_result = pickle.dumps(result)
            db.session.commit()
            
            # Sincroniza os novos grupos no banco de dados
            sincronizar_tabelas_dos_grupos(turma.id, result)
            
            flash('Grupos gerados com sucesso!', 'success')
            return redirect(url_for('view_turma', turma_id=turma_id))

    # --- LÓGICA GET ---
    base_data_info, groups_result_dict, ranking_data = None, None, None
    
    if turma.base_data:
        try:
            df = pd.read_csv(BytesIO(turma.base_data))
            alunos = sorted({str(a).replace('_', ' ') for a in df['aluno'].unique()})
            base_data_info = {"alunos": len(alunos), "dias": len(df['dia'].unique()), "horarios": len(df['horario'].unique()), "lista_alunos": alunos}
            ranking_data = calcular_ranking_disponibilidade(df.copy())
        except Exception as e:
            print(f"Erro ao ler base_data do banco: {e}")

    if turma.groups_result:
        try:
            groups_result_dict = pickle.loads(turma.groups_result)
            
            grupos_originais = groups_result_dict.get('grupos', [])
            slots_por_aluno = groups_result_dict.get('slots_por_aluno', {})
            grupos_para_template = []
            for group in grupos_originais:
                novo_grupo = group.copy()
                membros_com_sugestoes = []
                for membro_item in group.get('Membros', []):
                    nome_aluno = _get_nome_membro(membro_item)
                    sugs = sugerir_trocas(nome_aluno, grupos_originais, slots_por_aluno)
                    membros_com_sugestoes.append({'nome': nome_aluno.replace('_', ' '), 'sugestoes': sugs})
                novo_grupo['Membros'] = membros_com_sugestoes
                grupos_para_template.append(novo_grupo)
            groups_result_dict['grupos'] = grupos_para_template
        except Exception as e:
            print(f"Erro ao carregar groups_result do banco: {e}")

    # Selects dos blocos de imersão (janelas em hora cheia); a planilha usa HORARIOS_GRADE
    horarios_disponiveis = list(range(7, 19))
    nomes_blocos_opcoes = ["Manhã", "Tarde", "Noite", "Extra"]
    dias_semana_opcoes = [
        {"valor": "Seg", "texto": "Segunda-feira"},
        {"valor": "Ter", "texto": "Terça-feira"},
        {"valor": "Qua", "texto": "Quarta-feira"},
        {"valor": "Qui", "texto": "Quinta-feira"},
        {"valor": "Sex", "texto": "Sexta-feira"},
        {"valor": "Sab", "texto": "Sábado"}
    ]

    has_security_question = bool(user.security_question and user.security_answer_hash)

    return render_template('index.html', 
                           turma=turma,
                           base_data_info=base_data_info, 
                           groups_result=groups_result_dict, 
                           ranking_data=ranking_data,
                           horarios_disponiveis=horarios_disponiveis,
                           nomes_blocos_opcoes=nomes_blocos_opcoes,    
                           dias_semana_opcoes=dias_semana_opcoes,
                           has_security_question=has_security_question)

@app.route('/trocar_aluno/<int:turma_id>', methods=['POST'])
@login_required
def trocar_aluno(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.groups_result:
        flash('Dados de grupos não encontrados. Por favor, gere os grupos novamente.', 'error')
        return redirect(url_for('view_turma', turma_id=turma_id))

    try:
        result = pickle.loads(turma.groups_result)
    except Exception as e:
        flash('Erro ao carregar os dados dos grupos do banco de dados.', 'error')
        return redirect(url_for('view_turma', turma_id=turma_id))
    
    grupos = result['grupos']
    
    # 2. Obter dados do formulário de forma segura
    aluno_a = request.form.get('aluno_a')
    aluno_b = request.form.get('aluno_b') # Será None no caso de "mover"

    grupo_a_id_str = request.form.get('grupo_a_id')
    grupo_b_id_str = request.form.get('grupo_b_id')

    # Valida e converte os IDs
    grupo_a_id = int(grupo_a_id_str) if grupo_a_id_str else None
    grupo_b_id = int(grupo_b_id_str) if grupo_b_id_str else None

    # Verifica se os dados essenciais estão presentes
    if not aluno_a or not grupo_a_id or (aluno_b and not grupo_b_id) or (not aluno_b and not grupo_b_id):
        flash('Informação incompleta para realizar a movimentação.', 'error')
        return redirect(url_for('view_turma', turma_id=turma_id))

    # 3. Executar a lógica de troca/movimentação
    grupos_modificados = executar_troca(aluno_a, grupo_a_id, grupos, aluno_b, grupo_b_id)
    
    # 4. Salvar os novos dados no banco
    result['grupos'] = grupos_modificados
    turma.groups_result = pickle.dumps(result)
    db.session.commit()
    
    # Sincronizar tabelas relacionais de grupos e alunos
    sincronizar_tabelas_dos_grupos(turma.id, result)
        
    if aluno_b:
        flash(f'Troca entre {aluno_a} e {aluno_b} realizada com sucesso!', 'success')
    else:
        flash(f'Aluno {aluno_a} movido para o Grupo {grupo_b_id} com sucesso!', 'success')
        
    return redirect(url_for('view_turma', turma_id=turma_id))

@app.route('/clear/<int:turma_id>')
@login_required 
def clear_session(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if turma:
        turma.base_data = None
        turma.groups_result = None
        
        # Deletar todos os Alunos e Grupos desta turma do banco relacional
        Aluno.query.filter_by(turma_id=turma_id).delete()
        Grupo.query.filter_by(turma_id=turma_id).delete()
        
        db.session.commit()
        
    user_data_folder = get_user_data_folder()
    if user_data_folder and os.path.exists(user_data_folder):
        shutil.rmtree(user_data_folder)
    
    flash('Os dados foram limpos com sucesso!', 'success')
    return redirect(url_for('view_turma', turma_id=turma_id))

@app.route('/download/csv/<int:turma_id>')
@login_required
def download_csv(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma:
        return "Turma não encontrada.", 404
        
    if not turma.em_imersao and turma.base_data:
        # Se a turma não for de imersão e tiver PDFs processados, envia os dados brutos de alunos (base_data)
        buffer = BytesIO(turma.base_data)
        return send_file(buffer, mimetype='text/csv', as_attachment=True, download_name=f'dados_turma_{turma_id}.csv')
        
    if not turma.groups_result:
        return "Nenhum grupo gerado.", 404
        
    try:
        result = pickle.loads(turma.groups_result)
    except Exception as e:
        return "Erro ao carregar dados do banco.", 500
        
    df = grupos_para_dataframe(result['grupos'])
    buffer = BytesIO(df.to_csv(index=False, sep=';', encoding='utf-8-sig').encode('utf-8-sig'))
    return send_file(buffer, mimetype='text/csv', as_attachment=True, download_name=f'grupos_{turma.nome}.csv')

@app.route('/download/pdf/<int:turma_id>')
@login_required
def download_pdf(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.groups_result:
        return "Nenhum grupo gerado.", 404
        
    try:
        result = pickle.loads(turma.groups_result)
    except Exception as e:
        return "Erro ao carregar dados do banco.", 500
        
    pdf_buffer = montar_pdf_grupos(result['grupos'])
    return send_file(pdf_buffer, mimetype='application/pdf', as_attachment=True, download_name='grupos.pdf')

@app.route('/aluno/<nome_aluno>/<int:turma_id>')
@login_required
def get_aluno_schedule(nome_aluno, turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.base_data:
        return render_template_string("<h3>Erro</h3><p>Dados não encontrados no banco.</p>"), 404

    try:
        df = pd.read_csv(BytesIO(turma.base_data))
    except Exception as e:
        return render_template_string("<h3>Erro</h3><p>Erro ao ler dados do banco.</p>"), 500
    
    # Comparar o nome da URL (com espaços) com a coluna do df
    aluno_df = df[df['aluno'].str.replace('_', ' ').str.lower() == nome_aluno.lower()]

    if aluno_df.empty:
        return render_template_string(f"<h3>Aluno não encontrado</h3><p>O aluno '{nome_aluno}' não foi encontrado nos dados carregados.</p>"), 404
    
    # Pivoting (usando a coluna 'horario')
    schedule_pivot = aluno_df.pivot_table(
        index='horario', 
        columns='dia', 
        values='materia',
        aggfunc='first', 
        fill_value='---'
    ).sort_index()

    # Garantir que os dias da semana sejam exibidos corretamente
    dias_map = {
        'Seg': 'Segunda', 'Ter': 'Terça', 'Qua': 'Quarta',
        'Qui': 'Quinta', 'Sex': 'Sexta', 'Sab': 'Sábado'
    }
    
    for dia_abrev in dias_map.keys():
        if dia_abrev not in schedule_pivot.columns:
            schedule_pivot[dia_abrev] = '---'
            
    schedule_pivot.rename(columns=dias_map, inplace=True)
    schedule_pivot = schedule_pivot[list(dias_map.values())]

    return render_template("_aluno_schedule.html", aluno=nome_aluno, schedule=schedule_pivot)

@app.route('/download/schedule/xlsx/<int:turma_id>')
@login_required
def download_schedule_xlsx(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.base_data:
        return "Dados não encontrados.", 404
        
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        df = pd.read_csv(BytesIO(turma.base_data))
        
        # Filtrar por grupos selecionados (se houver parâmetro ?grupos=...)
        groups_param = request.args.get('grupos')
        if groups_param and groups_param != 'all' and turma.groups_result:
            try:
                result_g = pickle.loads(turma.groups_result)
                grupos_list = result_g.get('grupos', [])
                sem_grupo_list = result_g.get('alunos_sem_grupo', [])
                
                allowed_students = set()
                selected_ids = [x.strip() for x in groups_param.split(',')]
                for s_id in selected_ids:
                    if s_id == 'sem-grupo':
                        for a in sem_grupo_list:
                            allowed_students.add(a.replace('_', ' ').strip().lower())
                    else:
                        try:
                            g_num = int(s_id)
                            for g in grupos_list:
                                if int(g.get("Grupo", 0)) == g_num:
                                    for m in g.get("Membros", []):
                                        allowed_students.add(_get_nome_membro(m).replace('_', ' ').strip().lower())
                        except ValueError:
                            pass
                            
                if allowed_students:
                    df = df[df['aluno'].str.replace('_', ' ').str.strip().str.lower().isin(allowed_students)]
            except Exception as ex_filter:
                print(f"Erro ao filtrar por grupos: {ex_filter}")
        
        # Mapeamentos de dia
        dias_semana = ['Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado']
        dias_completos = {'Seg': 'Segunda', 'Ter': 'Terça', 'Qua': 'Quarta', 'Qui': 'Quinta', 'Sex': 'Sexta', 'Sab': 'Sábado'}
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Grade de Horários"
        ws.views.sheetView[0].showGridLines = True
        
        # Estilos
        font_header_dia = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        fill_header_dia = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid') # Primary color (indigo)
        
        font_header_hora = Font(name='Segoe UI', size=9, bold=True, color='334155')
        fill_header_hora = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid') # Slate 100
        
        font_aluno = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
        fill_aluno = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid') # Slate 50
        
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        
        thin_border_side = Side(border_style="thin", color="CBD5E1") # Slate 300
        border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        # Cores dos status
        STATUS_COLORS = {
            'L': 'D4EDDA',   # Free (green)
            'AC': 'FFE0B2',  # Acompanhamento Comportamental (orange)
            'PSI': 'EBD4FC', # Psicoterapia (purple)
            'X': 'FCC2C2',   # Aula de Graduação (red)
            'IM': 'CCE5FF',  # Imersão (blue)
            'AP': 'D1ECF1',  # Acompanhamento Pedagógico (mint)
            'R': 'E2E3E5',   # Realocado (grey)
        }
        
        font_status = Font(name='Segoe UI', size=9, bold=True, color='1E293B')
        
        # Cabeçalho 1 (Dias)
        ws.merge_cells('A1:A2')
        ws['A1'] = "Estudante"
        ws['A1'].font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        ws['A1'].fill = fill_header_dia
        ws['A1'].alignment = align_center
        ws['A1'].border = border_all
        ws['A2'].border = border_all
        
        col_idx = 2
        for dia in dias_semana:
            start_col = col_idx
            for hora in HORARIOS_GRADE:
                cell_hora = ws.cell(row=2, column=col_idx)
                cell_hora.value = hora
                cell_hora.font = font_header_hora
                cell_hora.fill = fill_header_hora
                cell_hora.alignment = align_center
                cell_hora.border = border_all
                col_idx += 1
                
            end_col = col_idx - 1
            start_letter = get_column_letter(start_col)
            end_letter = get_column_letter(end_col)
            ws.merge_cells(f"{start_letter}1:{end_letter}1")
            cell_dia = ws[f"{start_letter}1"]
            cell_dia.value = dia
            cell_dia.font = font_header_dia
            cell_dia.fill = fill_header_dia
            cell_dia.alignment = align_center
            
            for c in range(start_col, end_col + 1):
                ws.cell(row=1, column=c).border = border_all
                
        # Varrer alunos e preencher as células
        alunos_unicos = sorted(df['aluno'].unique())
        schedules = {}
        for aluno in alunos_unicos:
            nome_limpo = aluno.replace('_', ' ')
            schedules[nome_limpo] = {}
            for dia in dias_semana:
                for slot in HORARIOS_GRADE:
                    schedules[nome_limpo][f"{dia}_{slot}"] = 'L'
                    
        for _, row in df.iterrows():
            materia = row['materia']
            if pd.isna(materia) or str(materia).strip() in ('', '---'):
                continue
                
            intervalo = obter_intervalo_minutos(row['horario'])
            if intervalo is None:
                continue
                
            nome_limpo = str(row['aluno']).replace('_', ' ')
            dia_completo = dias_completos.get(row['dia'], row['dia'])
            materia_str = str(materia).strip()
            
            if materia_str in ('Acompanhamento Comportamental', 'AC') or 'Acompanhamento Comportamental' in materia_str:
                status = 'AC'
            elif materia_str in ('Psicoterapia', 'PSI') or 'Psicoterapia' in materia_str:
                status = 'PSI'
            elif materia_str in ('Imersão', 'IM') or 'Imersão' in materia_str:
                status = 'IM'
            elif materia_str in ('Acompanhamento Pedagógico', 'AP') or 'Acompanhamento Pedagógico' in materia_str:
                status = 'AP'
            elif materia_str in ('Realocado', 'R') or 'Realocado' in materia_str:
                status = 'R'
            elif materia_str.startswith('Custom:'):
                status = materia_str[7:].strip()
            else:
                status = 'X'
                
            ini, fim = intervalo
            aluno_sched = schedules.get(nome_limpo)
            if aluno_sched is not None:
                for slot in slots_ocupados_no_intervalo(ini, fim):
                    aluno_sched[f"{dia_completo}_{slot}"] = status

        row_idx = 3
        for aluno_nome in sorted(schedules.keys()):
            cell_aluno = ws.cell(row=row_idx, column=1)
            cell_aluno.value = aluno_nome
            cell_aluno.font = font_aluno
            cell_aluno.fill = fill_aluno
            cell_aluno.alignment = align_left
            cell_aluno.border = border_all
            
            col_idx = 2
            for dia in dias_semana:
                for slot in HORARIOS_GRADE:
                    cell = ws.cell(row=row_idx, column=col_idx)
                    status = schedules[aluno_nome][f"{dia}_{slot}"]
                    cell.value = status
                    cell.font = font_status
                    cell.alignment = align_center
                    cell.border = border_all
                    
                    color = STATUS_COLORS.get(status)
                    if color:
                        cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
                    else:
                        cell.fill = PatternFill(start_color='F5F3FF', end_color='F5F3FF', fill_type='solid')
                        
                    col_idx += 1
            row_idx += 1
            
        ws.column_dimensions['A'].width = 30
        for c in range(2, col_idx):
            col_letter = get_column_letter(c)
            ws.column_dimensions[col_letter].width = 6
            
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 20
        for r in range(3, row_idx):
            ws.row_dimensions[r].height = 24
            
        out_buf = BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        
        return send_file(
            out_buf, 
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 
            as_attachment=True, 
            download_name=f'grade_horarios_{turma.nome}.xlsx'
        )
    except Exception as e:
        print(f"Erro ao exportar XLSX de grade: {e}")
        return f"Erro ao gerar planilha Excel: {e}", 500

@app.route('/download/schedule/csv/<int:turma_id>')
@login_required
def download_schedule_csv(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.base_data:
        return "Dados não encontrados.", 404
        
    try:
        df = pd.read_csv(BytesIO(turma.base_data))
        
        # Filtrar por grupos selecionados (se houver parâmetro ?grupos=...)
        groups_param = request.args.get('grupos')
        if groups_param and groups_param != 'all' and turma.groups_result:
            try:
                result_g = pickle.loads(turma.groups_result)
                grupos_list = result_g.get('grupos', [])
                sem_grupo_list = result_g.get('alunos_sem_grupo', [])
                
                allowed_students = set()
                selected_ids = [x.strip() for x in groups_param.split(',')]
                for s_id in selected_ids:
                    if s_id == 'sem-grupo':
                        for a in sem_grupo_list:
                            allowed_students.add(a.replace('_', ' ').strip().lower())
                    else:
                        try:
                            g_num = int(s_id)
                            for g in grupos_list:
                                if int(g.get("Grupo", 0)) == g_num:
                                    for m in g.get("Membros", []):
                                        allowed_students.add(_get_nome_membro(m).replace('_', ' ').strip().lower())
                        except ValueError:
                            pass
                            
                if allowed_students:
                    df = df[df['aluno'].str.replace('_', ' ').str.strip().str.lower().isin(allowed_students)]
            except Exception as ex_filter:
                print(f"Erro ao filtrar por grupos: {ex_filter}")
        
        dias_semana = ['Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado']
        dias_completos = {'Seg': 'Segunda', 'Ter': 'Terça', 'Qua': 'Quarta', 'Qui': 'Quinta', 'Sex': 'Sexta', 'Sab': 'Sábado'}
        
        alunos_unicos = sorted(df['aluno'].unique())
        schedules = {}
        for aluno in alunos_unicos:
            nome_limpo = aluno.replace('_', ' ')
            schedules[nome_limpo] = {}
            for dia in dias_semana:
                for slot in HORARIOS_GRADE:
                    schedules[nome_limpo][f"{dia}_{slot}"] = 'L'
                    
        for _, row in df.iterrows():
            materia = row['materia']
            if pd.isna(materia) or str(materia).strip() in ('', '---'):
                continue
                
            intervalo = obter_intervalo_minutos(row['horario'])
            if intervalo is None:
                continue
                
            nome_limpo = str(row['aluno']).replace('_', ' ')
            dia_completo = dias_completos.get(row['dia'], row['dia'])
            materia_str = str(materia).strip()
            
            if materia_str in ('Acompanhamento Comportamental', 'AC') or 'Acompanhamento Comportamental' in materia_str:
                status = 'AC'
            elif materia_str in ('Psicoterapia', 'PSI') or 'Psicoterapia' in materia_str:
                status = 'PSI'
            elif materia_str in ('Imersão', 'IM') or 'Imersão' in materia_str:
                status = 'IM'
            elif materia_str in ('Acompanhamento Pedagógico', 'AP') or 'Acompanhamento Pedagógico' in materia_str:
                status = 'AP'
            elif materia_str in ('Realocado', 'R') or 'Realocado' in materia_str:
                status = 'R'
            elif materia_str.startswith('Custom:'):
                status = materia_str[7:].strip()
            else:
                status = 'X'
                
            ini, fim = intervalo
            aluno_sched = schedules.get(nome_limpo)
            if aluno_sched is not None:
                for slot in slots_ocupados_no_intervalo(ini, fim):
                    aluno_sched[f"{dia_completo}_{slot}"] = status
                    
        # Montar cabeçalhos
        headers = ["Estudante"]
        for dia in dias_semana:
            for slot in HORARIOS_GRADE:
                headers.append(f"{dia} - {slot}")
                
        # Montar linhas
        rows = []
        for aluno_nome in sorted(schedules.keys()):
            row = [aluno_nome]
            for dia in dias_semana:
                for slot in HORARIOS_GRADE:
                    row.append(schedules[aluno_nome][f"{dia}_{slot}"])
            rows.append(row)
            
        grid_df = pd.DataFrame(rows, columns=headers)
        buffer = BytesIO(grid_df.to_csv(index=False, sep=';', encoding='utf-8-sig').encode('utf-8-sig'))
        return send_file(
            buffer, 
            mimetype='text/csv', 
            as_attachment=True, 
            download_name=f'grade_horarios_{turma.nome}.csv'
        )
    except Exception as e:
        print(f"Erro ao exportar CSV de grade: {e}")
        return f"Erro ao gerar CSV de grade: {e}", 500

@app.route('/download/groups/xlsx/<int:turma_id>')
@login_required
def download_groups_xlsx(turma_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.groups_result:
        return "Nenhum grupo gerado.", 404
        
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        result = pickle.loads(turma.groups_result)
        grupos = result.get('grupos', [])
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Grupos Formados"
        ws.views.sheetView[0].showGridLines = True
        
        # Estilos
        font_title = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
        fill_title = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid') # Indigo
        
        font_group_header = Font(name='Segoe UI', size=11, bold=True, color='1E293B')
        fill_group_header = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid') # Slate 100
        
        font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
        fill_header = PatternFill(start_color='64748B', end_color='64748B', fill_type='solid') # Slate 500
        
        font_data = Font(name='Segoe UI', size=10, color='334155')
        
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        
        thin_side = Side(border_style="thin", color="CBD5E1")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        # Título
        ws.merge_cells('A1:D1')
        ws['A1'] = f"Grupos Formados — Turma: {turma.nome}"
        ws['A1'].font = font_title
        ws['A1'].fill = fill_title
        ws['A1'].alignment = align_center
        ws['A1'].border = border_all
        ws.row_dimensions[1].height = 36
        
        # Headers
        headers = ["Grupo", "Dia", "Horário", "Estudante"]
        for c_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=c_idx)
            cell.value = h
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_all
        ws.row_dimensions[2].height = 24
        
        row_idx = 3
        for g in sorted(grupos, key=lambda x: x.get("Grupo", 0)):
            membros_nomes = [_get_nome_membro(m) for m in g.get("Membros", [])]
            num_membros = len(membros_nomes)
            if num_membros == 0:
                continue
                
            start_row = row_idx
            for m_idx, m in enumerate(sorted(membros_nomes)):
                ws.cell(row=row_idx, column=4, value=m).font = font_data
                ws.cell(row=row_idx, column=4).border = border_all
                ws.cell(row=row_idx, column=4).alignment = align_left
                
                # Preencher temporariamente para caso não mescle ou para bordas
                ws.cell(row=row_idx, column=1, value=f"Grupo {g.get('Grupo')}").border = border_all
                ws.cell(row=row_idx, column=2, value=g.get('Dia')).border = border_all
                ws.cell(row=row_idx, column=3, value=g.get('Horário')).border = border_all
                
                row_idx += 1
            end_row = row_idx - 1
            
            ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
            ws.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
            
            ws.cell(row=start_row, column=1).font = font_group_header
            ws.cell(row=start_row, column=1).alignment = align_center
            
            ws.cell(row=start_row, column=2).font = font_data
            ws.cell(row=start_row, column=2).alignment = align_center
            
            ws.cell(row=start_row, column=3).font = font_data
            ws.cell(row=start_row, column=3).alignment = align_center
            
            for r in range(start_row, end_row + 1):
                ws.cell(row=r, column=1).fill = fill_group_header
                ws.cell(row=r, column=2).fill = fill_group_header
                ws.cell(row=r, column=3).fill = fill_group_header
                
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 30
        
        out_buf = BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        
        return send_file(
            out_buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'grupos_{turma.nome}.xlsx'
        )
    except Exception as e:
        print(f"Erro ao exportar grupos XLSX: {e}")
        return f"Erro ao gerar Excel de grupos: {e}", 500

@app.route('/download/grupo/<int:turma_id>/<int:grupo_id>')
@login_required
def download_grupo_individual(turma_id, grupo_id):
    user = User.query.filter_by(username=session['username']).first()
    turma = Turma.query.filter_by(id=turma_id, user_id=user.id).first()
    if not turma or not turma.groups_result:
        return "Nenhum grupo gerado.", 404
        
    try:
        result = pickle.loads(turma.groups_result)
        grupos = result.get('grupos', [])
        
        grupo_especifico = None
        for g in grupos:
            if int(g.get("Grupo", 0)) == int(grupo_id):
                grupo_especifico = g
                break
                
        if not grupo_especifico:
            return "Grupo não encontrado.", 404
            
        membros_nomes = [_get_nome_membro(m) for m in grupo_especifico.get("Membros", [])]
        export_format = request.args.get('format', 'csv')
        
        if export_format == 'xlsx':
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"Grupo {grupo_id}"
            ws.views.sheetView[0].showGridLines = True
            
            # Estilos
            font_title = Font(name='Segoe UI', size=13, bold=True, color='FFFFFF')
            fill_title = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid') # Indigo
            
            font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
            fill_header = PatternFill(start_color='64748B', end_color='64748B', fill_type='solid') # Slate 500
            
            font_data = Font(name='Segoe UI', size=10, color='334155')
            
            align_center = Alignment(horizontal='center', vertical='center')
            align_left = Alignment(horizontal='left', vertical='center')
            
            thin_side = Side(border_style="thin", color="CBD5E1")
            border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
            
            # Título
            ws.merge_cells('A1:C1')
            ws['A1'] = f"Grupo {grupo_id} — Turma: {turma.nome}"
            ws['A1'].font = font_title
            ws['A1'].fill = fill_title
            ws['A1'].alignment = align_center
            ws['A1'].border = border_all
            ws.row_dimensions[1].height = 32
            
            # Informações do grupo
            ws.merge_cells('A2:C2')
            ws['A2'] = f"Dia: {grupo_especifico.get('Dia')} | Horário: {grupo_especifico.get('Horário')}"
            ws['A2'].font = Font(name='Segoe UI', size=10, bold=True, color='334155')
            ws['A2'].fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
            ws['A2'].alignment = align_center
            ws['A2'].border = border_all
            ws.row_dimensions[2].height = 24
            
            # Headers
            headers = ["Grupo", "Estudante", "Status"]
            for c_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=3, column=c_idx)
                cell.value = h
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = align_center
                cell.border = border_all
            ws.row_dimensions[3].height = 24
            
            row_idx = 4
            for m in sorted(membros_nomes):
                ws.cell(row=row_idx, column=1, value=f"Grupo {grupo_id}").font = font_data
                ws.cell(row=row_idx, column=1).alignment = align_center
                ws.cell(row=row_idx, column=1).border = border_all
                
                ws.cell(row=row_idx, column=2, value=m).font = font_data
                ws.cell(row=row_idx, column=2).alignment = align_left
                ws.cell(row=row_idx, column=2).border = border_all
                
                ws.cell(row=row_idx, column=3, value="Membro").font = font_data
                ws.cell(row=row_idx, column=3).alignment = align_center
                ws.cell(row=row_idx, column=3).border = border_all
                
                ws.row_dimensions[row_idx].height = 20
                row_idx += 1
                
            ws.column_dimensions['A'].width = 15
            ws.column_dimensions['B'].width = 30
            ws.column_dimensions['C'].width = 15
            
            out_buf = BytesIO()
            wb.save(out_buf)
            out_buf.seek(0)
            
            return send_file(
                out_buf,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                as_attachment=True,
                download_name=f'grupo_{grupo_id}_{turma.nome}.xlsx'
            )
        else:
            linhas = []
            for m in sorted(membros_nomes):
                linhas.append({
                    "Grupo": f"Grupo {grupo_id}", 
                    "Dia": grupo_especifico.get("Dia"), 
                    "Horário": grupo_especifico.get("Horário"), 
                    "Aluno": m
                })
            group_df = pd.DataFrame(linhas)
            buffer = BytesIO(group_df.to_csv(index=False, sep=';', encoding='utf-8-sig').encode('utf-8-sig'))
            return send_file(
                buffer, 
                mimetype='text/csv', 
                as_attachment=True, 
                download_name=f'grupo_{grupo_id}_{turma.nome}.csv'
            )
    except Exception as e:
        print(f"Erro ao exportar grupo individual: {e}")
        return f"Erro ao exportar grupo: {e}", 500

@app.route('/download/all_schedules/<string:export_format>')
@login_required
def download_all_schedules(export_format):
    user = User.query.filter_by(username=session['username']).first()
    if not user:
        return "Usuário não encontrado.", 404
        
    search_val = request.args.get('search', '').strip().lower()
    turmas_param = request.args.get('turmas', '').strip()
    students_param = request.args.get('students', '').strip()
    
    query = Turma.query.filter_by(user_id=user.id)
    if turmas_param:
        names_list = [x.strip() for x in turmas_param.split(',')]
        if names_list:
            query = query.filter(Turma.nome.in_(names_list))
    turmas = query.all()
    
    allowed_students_filter = None
    if students_param:
        allowed_students_filter = {x.strip().lower() for x in students_param.split(',')}
        
    dias_semana = ['Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado']
    dias_completos = {'Seg': 'Segunda', 'Ter': 'Terça', 'Qua': 'Quarta', 'Qui': 'Quinta', 'Sex': 'Sexta', 'Sab': 'Sábado'}
    
    schedules = {}
    
    for t in turmas:
        if not t.base_data:
            continue
        try:
            df = pd.read_csv(BytesIO(t.base_data))
            alunos_unicos = df['aluno'].unique()
            
            for aluno in alunos_unicos:
                nome_limpo = aluno.replace('_', ' ')
                
                if search_val and search_val not in nome_limpo.lower():
                    continue
                    
                if allowed_students_filter and nome_limpo.lower() not in allowed_students_filter:
                    continue
                    
                schedules[nome_limpo] = {
                    'turma': t.nome,
                }
                for dia in dias_semana:
                    for slot in HORARIOS_GRADE:
                        schedules[nome_limpo][f"{dia}_{slot}"] = 'L'
                        
            for _, row in df.iterrows():
                materia = row['materia']
                if pd.isna(materia) or str(materia).strip() in ('', '---'):
                    continue
                intervalo = obter_intervalo_minutos(row['horario'])
                if intervalo is None:
                    continue
                
                nome_limpo = str(row['aluno']).replace('_', ' ')
                if nome_limpo not in schedules:
                    continue
                    
                dia_completo = dias_completos.get(row['dia'], row['dia'])
                materia_str = str(materia).strip()
                
                if materia_str in ('Acompanhamento Comportamental', 'AC') or 'Acompanhamento Comportamental' in materia_str:
                    status = 'AC'
                elif materia_str in ('Psicoterapia', 'PSI') or 'Psicoterapia' in materia_str:
                    status = 'PSI'
                elif materia_str in ('Imersão', 'IM') or 'Imersão' in materia_str:
                    status = 'IM'
                elif materia_str in ('Acompanhamento Pedagógico', 'AP') or 'Acompanhamento Pedagógico' in materia_str:
                    status = 'AP'
                elif materia_str in ('Realocado', 'R') or 'Realocado' in materia_str:
                    status = 'R'
                elif materia_str.startswith('Custom:'):
                    status = materia_str[7:].strip()
                else:
                    status = 'X'
                    
                ini, fim = intervalo
                for slot in slots_ocupados_no_intervalo(ini, fim):
                    schedules[nome_limpo][f"{dia_completo}_{slot}"] = status
        except Exception as e:
            print(f"Erro ao processar turma {t.id} no download geral: {e}")
            
    if not schedules:
        return "Nenhum aluno encontrado com os filtros aplicados.", 404
        
    if export_format == 'xlsx':
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Grade Geral"
        ws.views.sheetView[0].showGridLines = True
        
        # Estilos
        font_header_dia = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        fill_header_dia = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        
        font_header_hora = Font(name='Segoe UI', size=9, bold=True, color='334155')
        fill_header_hora = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
        
        font_aluno = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
        fill_aluno = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        
        font_turma = Font(name='Segoe UI', size=9, color='475569')
        fill_turma = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        
        thin_border_side = Side(border_style="thin", color="CBD5E1")
        border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        STATUS_COLORS = {
            'L': 'D4EDDA',
            'AC': 'FFE0B2',
            'PSI': 'EBD4FC',
            'X': 'FCC2C2',
            'IM': 'CCE5FF',
            'AP': 'D1ECF1',
            'R': 'E2E3E5',
        }
        font_status = Font(name='Segoe UI', size=9, bold=True, color='1E293B')
        
        # Cabeçalho
        ws.merge_cells('A1:A2')
        ws['A1'] = "Estudante"
        ws['A1'].font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        ws['A1'].fill = fill_header_dia
        ws['A1'].alignment = align_center
        ws['A1'].border = border_all
        ws['A2'].border = border_all
        
        ws.merge_cells('B1:B2')
        ws['B1'] = "Turma"
        ws['B1'].font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        ws['B1'].fill = fill_header_dia
        ws['B1'].alignment = align_center
        ws['B1'].border = border_all
        ws['B2'].border = border_all
        
        col_idx = 3
        for dia in dias_semana:
            start_col = col_idx
            for hora in HORARIOS_GRADE:
                cell_hora = ws.cell(row=2, column=col_idx)
                cell_hora.value = hora
                cell_hora.font = font_header_hora
                cell_hora.fill = fill_header_hora
                cell_hora.alignment = align_center
                cell_hora.border = border_all
                col_idx += 1
                
            end_col = col_idx - 1
            start_letter = get_column_letter(start_col)
            end_letter = get_column_letter(end_col)
            ws.merge_cells(f"{start_letter}1:{end_letter}1")
            cell_dia = ws[f"{start_letter}1"]
            cell_dia.value = dia
            cell_dia.font = font_header_dia
            cell_dia.fill = fill_header_dia
            cell_dia.alignment = align_center
            
            for c in range(start_col, end_col + 1):
                ws.cell(row=1, column=c).border = border_all
                
        row_idx = 3
        for aluno_nome in sorted(schedules.keys()):
            cell_aluno = ws.cell(row=row_idx, column=1)
            cell_aluno.value = aluno_nome
            cell_aluno.font = font_aluno
            cell_aluno.fill = fill_aluno
            cell_aluno.alignment = align_left
            cell_aluno.border = border_all
            
            cell_turma = ws.cell(row=row_idx, column=2)
            cell_turma.value = schedules[aluno_nome]['turma']
            cell_turma.font = font_turma
            cell_turma.fill = fill_turma
            cell_turma.alignment = align_left
            cell_turma.border = border_all
            
            col_idx = 3
            for dia in dias_semana:
                for slot in HORARIOS_GRADE:
                    cell = ws.cell(row=row_idx, column=col_idx)
                    status = schedules[aluno_nome][f"{dia}_{slot}"]
                    cell.value = status
                    cell.font = font_status
                    cell.alignment = align_center
                    cell.border = border_all
                    
                    color = STATUS_COLORS.get(status)
                    if color:
                        cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
                    else:
                        cell.fill = PatternFill(start_color='F5F3FF', end_color='F5F3FF', fill_type='solid')
                    col_idx += 1
            row_idx += 1
            
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 15
        for c in range(3, col_idx):
            col_letter = get_column_letter(c)
            ws.column_dimensions[col_letter].width = 6
            
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 20
        for r in range(3, row_idx):
            ws.row_dimensions[r].height = 24
            
        out_buf = BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        
        return send_file(
            out_buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='grade_geral_alunos.xlsx'
        )
    else:
        headers = ["Estudante", "Turma"]
        for dia in dias_semana:
            for slot in HORARIOS_GRADE:
                headers.append(f"{dia} - {slot}")
                
        rows = []
        for aluno_nome in sorted(schedules.keys()):
            row = [aluno_nome, schedules[aluno_nome]['turma']]
            for dia in dias_semana:
                for slot in HORARIOS_GRADE:
                    row.append(schedules[aluno_nome][f"{dia}_{slot}"])
            rows.append(row)
            
        grid_df = pd.DataFrame(rows, columns=headers)
        buffer = BytesIO(grid_df.to_csv(index=False, sep=';', encoding='utf-8-sig').encode('utf-8-sig'))
        return send_file(
            buffer,
            mimetype='text/csv',
            as_attachment=True,
            download_name='grade_geral_alunos.csv'
        )

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=int(os.environ.get('PORT', 5000)))
